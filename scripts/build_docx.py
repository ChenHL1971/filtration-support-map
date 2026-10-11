import re, html, xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src = open("/tmp/claude-0/-home-claude/cdd709eb-e6fa-5221-bfa3-4ba0261c0645/scratchpad/doc.xml").read()
root = ET.fromstring(src)
d = Document()
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(6)
for h in ["Heading 1", "Heading 2", "Heading 3", "Title"]:
    s = d.styles[h]; s.font.name = "Times New Roman"; s.font.color.rgb = None
    s.paragraph_format.keep_with_next = False; s.paragraph_format.keep_together = False      # no black marks beside headings
d.styles["Heading 1"].font.size = Pt(14); d.styles["Heading 2"].font.size = Pt(13); d.styles["Heading 3"].font.size = Pt(12)
sec = d.sections[0]; sec.left_margin = sec.right_margin = Inches(1)
# page number footer
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for t, txt in [("begin", None), (None, "PAGE"), ("end", None)]:
    if t: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
    else: e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
    r._r.append(e)

def runs(p_el, para):
    for t in p_el.iter("text"):
        b = t.find("bold")
        txt = (b.text if b is not None else t.text) or ""
        rr = para.add_run(txt); rr.bold = b is not None

section = None
NOINDENT = {"Title page", "Abstract", "References", "Tables", "Figure legends", "Supplementary tables"}
def add_par(el, style=None, indent=True):
    p = d.add_paragraph(style=style); runs(el, p)
    if indent and style is None: p.paragraph_format.first_line_indent = Inches(0.5)
    return p

kids = list(root)
for i, el in enumerate(kids):
    if el.tag == "paragraph":
        h = el.get("heading")
        text = "".join((t.find("bold").text if t.find("bold") is not None else t.text) or "" for t in el.iter("text"))
        if h == "1":
            p = d.add_paragraph(); rr = p.add_run(text); rr.bold = True; rr.font.size = Pt(16)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; continue
        if h == "2":
            section = text
            if text == "Title page": continue
            if text in ("Tables", "Figure legends", "Supplementary tables", "References"): d.add_page_break()
            if text in ("Introduction",): d.add_page_break()
            d.add_heading(text, level=1); continue
        if h == "3":
            d.add_heading(text, level=2); continue
        if not text.strip(): continue
        bold_lead = el.find("text/bold") is not None and list(el.iter("text"))[0].find("bold") is not None
        indent = section not in NOINDENT and not bold_lead and section is not None
        if section is None:   # draft line under title
            p = add_par(el, indent=False); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; continue
        add_par(el, indent=indent)
    elif el.tag == "list":
        ordered = el.get("kind") == "ordered"
        for n, li in enumerate(el.findall("listItem"), 1):
            for pe in li.findall("paragraph"):
                if section == "References":
                    p = d.add_paragraph(); p.add_run(f"{n}. "); runs(pe, p)
                    p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.first_line_indent = Inches(-0.35)
                else:
                    p = d.add_paragraph(style="List Number" if ordered else "List Bullet"); runs(pe, p)
    elif el.tag == "table":
        rows = el.findall("row"); ncol = len(rows[0].findall("cell"))
        t = d.add_table(rows=len(rows), cols=ncol); t.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row.findall("cell")):
                c = t.cell(ri, ci); c.paragraphs[0].text = ""
                ps = cell.findall("paragraph")
                for k, pe in enumerate(ps):
                    para = c.paragraphs[0] if k == 0 else c.add_paragraph()
                    runs(pe, para)
                    para.paragraph_format.line_spacing = 1.0; para.paragraph_format.space_after = Pt(0)
                    for rr in para.runs:
                        rr.font.size = Pt(9)
                        if cell.get("header") == "true": rr.bold = True
        d.add_paragraph()

# figures for review
d.add_page_break(); d.add_heading("Figures", level=1)
for f, cap in [("Fig1_study_design", "Figure 1"), ("FigNew_map", "Figure 2"), ("Fig2_fs_titan", "Figure 3"), ("Fig3_incremental_value", "Figure 4"), ("Fig4_finngen", "Figure 5")]:
    d.add_paragraph().add_run(cap).bold = True
    d.add_picture(f"/home/claude/audit/figs/{f}.png", width=Inches(6.3))
out = "/home/claude/audit/manuscript/Manuscript_Draft_FiltrationSupportMap_2026-10-11.docx"
d.save(out); print(out)
