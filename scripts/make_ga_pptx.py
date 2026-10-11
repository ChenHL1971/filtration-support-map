"""Editable PowerPoint graphical abstract (4 x 4 in -> 1200 x 1200 px at 300 dpi), native shapes only."""
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
BLUE, BLUE_BG = RGBColor(0x1F, 0x6F, 0xC5), RGBColor(0xE8, 0xF1, 0xFC)
ORANGE, ORANGE_BG = RGBColor(0xC9, 0x54, 0x1F), RGBColor(0xFD, 0xEE, 0xE7)
INK, MUTED, GREY_BG, GREY_LN = RGBColor(0x22, 0x22, 0x22), RGBColor(0x6B, 0x6B, 0x66), RGBColor(0xF2, 0xF2, 0xEF), RGBColor(0xA9, 0xA8, 0xA0)
prs = Presentation(); prs.slide_width = prs.slide_height = Inches(4)
s = prs.slides.add_slide(prs.slide_layouts[6])

def text(x, y, w, h, runs, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
    first = True
    for line in runs:                      # line = list of (text, size, bold, italic, color)
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False; p.alignment = align
        for t, size, bold, italic, color in line:
            r = p.add_run(); r.text = t; f = r.font; f.name = "Arial"; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color
    return tb
def box(x, y, w, h, fill, line, label=None, size=9, color=INK, bold=True):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.18; sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.color.rgb = line; sh.line.width = Pt(1.25); sh.shadow.inherit = False
    if label:
        tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, Inches(0.03))
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = label
        r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return sh
def arrow(x0, y0, x1, y1, color=MUTED, w=1.4):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = color; c.line.width = Pt(w)
    ln = c.line._get_or_add_ln(); te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c

# kidney (freeform bean) + title
cx, cy, sc = 0.42, 0.40, 0.23
pts = []
for k in range(0, 360, 6):
    t = math.radians(k); a = (t + math.pi) % (2 * math.pi) - math.pi
    r = 1 - 0.42 * math.exp(-(a / 0.55) ** 2)                               # concave hilum on the right
    pts.append((cx + sc * 0.80 * r * math.cos(t), cy - sc * 1.05 * math.sin(t)))
ff = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
ff.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=True)
kid = ff.convert_to_shape(); kid.fill.solid(); kid.fill.fore_color.rgb = RGBColor(0xE9, 0xA3, 0xA0); kid.line.color.rgb = RGBColor(0x9C, 0x4B, 0x47); kid.line.width = Pt(1); kid.shadow.inherit = False
text(0.78, 0.22, 3.1, 0.36, [[("Kidney function genetic signals", 12, True, False, INK)]], align=PP_ALIGN.LEFT)

# one fine-mapped signal -> two eGFR markers
d = s.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(1.89), Inches(0.72), Inches(0.22), Inches(0.22))
d.fill.solid(); d.fill.fore_color.rgb = INK; d.line.fill.background(); d.shadow.inherit = False
text(2.16, 0.70, 1.5, 0.26, [[("Fine-mapped signal", 8, False, False, MUTED)]], align=PP_ALIGN.LEFT)
arrow(1.92, 0.94, 1.15, 1.16, ORANGE); arrow(2.08, 0.94, 2.85, 1.16, BLUE)
box(0.30, 1.18, 1.55, 0.36, ORANGE_BG, ORANGE, "Creatinine-based eGFR", 8.5, ORANGE)
box(2.15, 1.18, 1.55, 0.36, BLUE_BG, BLUE, "Cystatin C-based eGFR", 8.5, BLUE)
arrow(1.08, 1.54, 1.55, 1.74); arrow(2.92, 1.54, 2.45, 1.74)
box(1.05, 1.76, 1.90, 0.32, GREY_BG, GREY_LN, "Signal-level colocalization", 8.5, INK)
arrow(2.0, 2.08, 2.0, 2.20)
text(0.6, 2.20, 2.8, 0.20, [[("Filtration-support annotation", 9, True, False, INK)]])
bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(2.43), Inches(2.5), Inches(0.13))
bar.adjustments[0] = 0.5; bar.line.fill.background(); bar.shadow.inherit = False
bar.fill.gradient(); bar.fill.gradient_angle = 0
bar.fill.gradient_stops[0].color.rgb = ORANGE; bar.fill.gradient_stops[1].color.rgb = BLUE
text(0.22, 2.40, 0.5, 0.2, [[("low", 8, False, True, MUTED)]], align=PP_ALIGN.RIGHT)
text(3.30, 2.40, 0.5, 0.2, [[("high", 8, False, True, MUTED)]], align=PP_ALIGN.LEFT)

# two consequences
arrow(1.55, 2.60, 1.05, 2.84, BLUE); arrow(2.45, 2.60, 2.95, 2.84, ORANGE)
box(0.15, 2.86, 1.78, 0.74, BLUE_BG, BLUE)
text(0.2, 2.92, 1.68, 0.64, [[("Signal interpretation", 9.5, True, False, BLUE)], [("Flags filtration-consistent signals", 8.5, False, False, INK)]])
box(2.07, 2.86, 1.78, 0.74, ORANGE_BG, ORANGE)
text(2.12, 2.92, 1.68, 0.64, [[("Gene prioritization", 9.5, True, False, ORANGE)], [("No meaningful gain in the tested model", 8.5, False, False, INK)]])
text(0.15, 3.66, 3.7, 0.28, [[("Informs signal interpretation, not effector-gene identification", 8.5, False, True, MUTED)]])
# drop theme style references (removes inherited shadows in every renderer)
for sh in s.shapes:
    st = sh._element.find(qn("p:style"))
    if st is not None: sh._element.remove(st)
prs.save("figs/ga/Graphical_Abstract.pptx"); print("saved")
