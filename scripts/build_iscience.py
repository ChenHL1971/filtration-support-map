"""Build iScience-formatted manuscript (Cell Press structure) from the master doc XML."""
import re, xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

root = ET.fromstring(open("/tmp/claude-0/-home-claude/cdd709eb-e6fa-5221-bfa3-4ba0261c0645/scratchpad/doc.xml").read())
REPL = [("as described in Methods", "as described in STAR Methods"), ("Supplementary Tables S1", "Tables S1"), ("Supplementary Table S1", "Table S1"), ("Supplementary Table S2", "Table S2"),
        ("Supplementary Table S3", "Table S3"), ("Supplementary Data 1", "Data S1"),
        ("There are also limitations. ", "")]
def fix(t):
    for a, b in REPL: t = t.replace(a, b)
    return t
GENES = ["SLC22A2", "SLC6A13", "SHROOM3", "TMEM60", "BCAS3", "WDR37", "WDR72", "CPS1", "UMOD"]
CIT = re.compile(r"\s?\[(\d+(?:\s?[–,\-]\s?\d+)*)\]")
TOK = re.compile(r"(\s?\[\d+(?:\s?[–,\-]\s?\d+)*\]|\b(?:" + "|".join(GENES) + r")\b)")
import json as _json, os as _os
CITEMAP = {int(k): v for k, v in _json.load(open(_os.environ["CITEMAP"])).items()} if _os.environ.get("CITEMAP") else {}
SEEN = []          # citation numbers (original numbering) in order of appearance
def remap(g):
    nums = []
    for part in g.replace(" ", "").replace("–", "-").split(","):
        if "-" in part: a, b = part.split("-"); nums += list(range(int(a), int(b) + 1))
        else: nums.append(int(part))
    for n in nums:
        if n not in SEEN: SEEN.append(n)
    nums = sorted(CITEMAP.get(n, n) for n in nums); out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ",".join(map(str, nums[i:j + 1]))); i = j + 1
    return ",".join(out)
def add_rich(p, txt, bold=False, size=None, italic=False):
    """Add text to paragraph p: [n] citations become superscript numbers, gene symbols italic."""
    for part in TOK.split(txt):
        if not part: continue
        m = CIT.fullmatch(part)
        if m: r = p.add_run(remap(m.group(1))); r.font.superscript = True
        else:
            r = p.add_run(part); r.italic = italic or part in GENES
        r.bold = bold
        if size: r.font.size = size
    return p
def text_of(el): return "".join(((t.find("bold").text if t.find("bold") is not None else t.text) or "") for t in el.iter("text"))

# ---- collect sections from master ----
sec, cur3 = None, None
S = {}            # h2 -> list of (kind, element, h3)
for el in root:
    if el.tag == "paragraph" and el.get("heading") == "2":
        sec = text_of(el); S.setdefault(sec, []); cur3 = None; continue
    if sec is None: continue
    if el.tag == "paragraph" and el.get("heading") == "3":
        cur3 = text_of(el); S[sec].append(("h3", el, cur3)); continue
    S[sec].append((el.tag, el, cur3))
title = text_of(root[0])
MASTER_REFS = []
for kind, el, h3 in S["References"]:
    if kind == "list":
        for li in el.findall("listItem"):
            for pe in li.findall("paragraph"): MASTER_REFS.append("".join((t.text or "") for t in pe.iter("text")))
    elif kind == "paragraph" and text_of(el).strip(): MASTER_REFS.append(text_of(el))
CODE_N = len(MASTER_REFS) + 1

import os
SUBMISSION = os.environ.get("SUBMISSION") == "1"
CODE_DOI = "10.5281/zenodo.23294824"
GITHUB = "https://github.com/ChenHL1971/filtration-support-map"
CODE_REF = "Chen, H.-L., Jeng, M.-J., Tang, Y.-H., Lin, C.-H., Teng, C.-T., and Wang, H.-H. (2026). Filtration-support map of kidney function genetic signals: code and data (version 1.1.0). Zenodo. https://doi.org/" + CODE_DOI
HIGHLIGHTS = ["A filtration-support map scores 1,681 of 2,051 fine-mapped eGFRcrea signals",
              "Support is lower at creatinine-specific than at GFR-associated loci",
              "Filtration support yields no meaningful gain in the tested gene-ranking framework",
              "Filtration-consistent signals do not directly identify causal genes"]
def new_doc(line_numbers=True, page_numbers=True):
    global d
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.line_spacing = 2.0; st.paragraph_format.space_after = Pt(0)
    for h in ["Heading 1", "Heading 2", "Heading 3"]:
        s = d.styles[h]; s.font.name = "Times New Roman"; s.font.color.rgb = None; s.paragraph_format.keep_with_next = False; s.paragraph_format.keep_together = False
    d.styles["Heading 1"].font.size = Pt(12); d.styles["Heading 2"].font.size = Pt(12); d.styles["Heading 3"].font.size = Pt(12)
    sec0 = d.sections[0]; sec0.left_margin = sec0.right_margin = Inches(1)
    fp = sec0.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = fp.add_run()
    for t, txt in ([("begin", None), (None, "PAGE"), ("end", None)] if page_numbers else []):
        if t: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
        else: e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)
    # line numbering
    if line_numbers:
        ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); sec0._sectPr.append(ln)

    return d
new_doc()

def H1(t): d.add_heading(t.upper(), level=1)
def H2(t): d.add_heading(t, level=2)
def P(t, indent=True, bold_lead=None, italic=False):
    p = d.add_paragraph()
    if bold_lead: p.add_run(bold_lead).bold = True
    add_rich(p, fix(t), italic=italic)
    if indent: p.paragraph_format.first_line_indent = Inches(0.5)
    return p
def P_el(el, indent=True):
    p = d.add_paragraph()
    groups = []   # merge adjacent text nodes with the same bold state so split citations ("[" + "3]") are recognised
    for t in el.iter("text"):
        b = t.find("bold"); txt = (b.text if b is not None else t.text) or ""
        if groups and groups[-1][1] == (b is not None): groups[-1][0] += txt
        else: groups.append([txt, b is not None])
    for txt, isb in groups: add_rich(p, fix(txt), bold=isb)
    lead_bold = el.find("text/bold") is not None and list(el.iter("text"))[0].find("bold") is not None
    if indent and not lead_bold: p.paragraph_format.first_line_indent = Inches(0.5)
    return p
def TABLE(el):
    rows = el.findall("row"); nc = len(rows[0].findall("cell"))
    t = d.add_table(rows=len(rows), cols=nc); t.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row.findall("cell")):
            c = t.cell(ri, ci); c.paragraphs[0].text = ""
            for k, pe in enumerate(cell.findall("paragraph")):
                para = c.paragraphs[0] if k == 0 else c.add_paragraph()
                groups = []
                for tx in pe.iter("text"):
                    b = tx.find("bold"); txt = (b.text if b is not None else tx.text) or ""
                    if groups and groups[-1][1] == (b is not None): groups[-1][0] += txt
                    else: groups.append([txt, b is not None])
                for txt, isb in groups: add_rich(para, fix(txt), bold=isb or cell.get("header") == "true", size=Pt(9))
                para.paragraph_format.line_spacing = 1.0
    d.add_paragraph()
def SIMPLE_TABLE(rows, header=True, size=9):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, r_ in enumerate(rows):
        for j, v in enumerate(r_):
            c = t.cell(i, j); c.paragraphs[0].text = ""; add_rich(c.paragraphs[0], v, bold=header and i == 0, size=Pt(size))
            c.paragraphs[0].paragraph_format.line_spacing = 1.0
    d.add_paragraph()
def emit(items, skip_h3=(), h3_level=2, stop_at=None, drop=lambda e: False):
    for kind, el, h3 in items:
        if h3 in skip_h3: continue
        if kind == "h3":
            if stop_at and text_of(el) == stop_at: break
            d.add_heading(text_of(el), level=h3_level); continue
        if kind == "paragraph":
            if not text_of(el).strip() or drop(el): continue
            P_el(el)
        elif kind == "table": TABLE(el)
        elif kind == "list":
            for li in el.findall("listItem"):
                for pe in li.findall("paragraph"):
                    p = d.add_paragraph(style="List Bullet")
                    for tx in pe.iter("text"): add_rich(p, fix(tx.text or ""))

# ---- title page ----
p = d.add_paragraph(); rr = p.add_run(title); rr.bold = True; rr.font.size = Pt(14); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
AUTHORS = [("Hui-Lan Chen", "1,2,3"), ("Mei-Jy Jeng", "1,4,5"), ("Yi-Hsuan Tang", "1,3"), ("Chien-Hung Lin", "3,6"),
           ("Chao-Ting Teng", "3"), ("Hsin-Hui Wang", "1,3,6,7,*")]
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, (nm, sup) in enumerate(AUTHORS):
    p.add_run(nm); r = p.add_run(sup); r.font.superscript = True
    if i == len(AUTHORS) - 2: p.add_run(", and ")
    elif i < len(AUTHORS) - 2: p.add_run(", ")
AFFS = ["Institute of Emergency and Critical Care Medicine, School of Medicine, National Yang Ming Chiao Tung University, Taipei 112, Taiwan",
        "Department of Pediatrics, Keelung Hospital, Ministry of Health and Welfare, Keelung 201, Taiwan",
        "Division of Pediatric Immunology and Nephrology, Department of Pediatrics, Taipei Veterans General Hospital, Taipei 112, Taiwan",
        "Department of Pediatrics, Taipei Veterans General Hospital, Taipei 112, Taiwan",
        "Department of Pediatrics, School of Medicine, National Yang Ming Chiao Tung University, Taipei 112, Taiwan",
        "Department of Pediatrics, Faculty of Medicine, School of Medicine, National Yang Ming Chiao Tung University, Taipei 112, Taiwan",
        "Lead contact"]
for k, a in enumerate(AFFS, 1):
    p = d.add_paragraph(); r = p.add_run(str(k)); r.font.superscript = True; p.add_run(a)
p = d.add_paragraph(); p.add_run("*Correspondence: hhwang@vghtpe.gov.tw")
P("ORCID: Hui-Lan Chen 0000-0001-8066-3338; Mei-Jy Jeng 0000-0002-7427-0367; Yi-Hsuan Tang 0000-0003-4240-5017; Chien-Hung Lin 0000-0003-0811-6529; Hsin-Hui Wang 0000-0003-4476-841x", indent=False)
H1("Summary")
SUMMARY = ("Genetic associations with creatinine-based estimated glomerular filtration rate (eGFRcrea) may reflect filtration or non-filtration determinants of creatinine. "
 "We developed a filtration-support map for 2,051 fine-mapped eGFRcrea signals using directionally concordant colocalization with cystatin C-based eGFR. "
 "Of 1,681 classifiable signals, 287 were filtration-consistent. "
 "Support was lower at previously classified creatinine-specific than GFR-associated loci (median 0.002 versus 0.392), indicating methodological consistency rather than independent validation. "
 "In the primary scorecard-based analysis, adding filtration support did not improve recovery of established kidney disease genes (difference in area under the precision\u2013recall curve, \u22120.0003; 95% CI \u22120.007 to 0.005). "
 "Filtration-consistent signals showed no statistically detectable enrichment for dialysis associations in FinnGen, although precision was limited. "
 "This public resource distinguishes cross-biomarker support for genetic signals from evidence identifying their effector genes.")
P(SUMMARY, indent=False)
d.add_page_break()

H1("Introduction"); emit(S["Introduction"])
H1("Results"); emit(S["Results"])
disc = S["Discussion"]
lim = [x for x in disc if x[0] == "paragraph" and text_of(x[1]).startswith("There are also limitations")]
H1("Discussion")
emit(disc, drop=lambda e: text_of(e).startswith("There are also limitations"), h3_level=2)
# remove the 'Conclusions' sub-heading look: handled by emit as heading level 2 (kept)
# ---- back matter in the order required by iScience final-file requirements ----
H1("Resource availability")
H2("Lead contact")
P("Requests for further information and resources should be directed to and will be fulfilled by the lead contact, Hsin-Hui Wang (hhwang@vghtpe.gov.tw).", indent=False)
H2("Materials availability")
P("This study did not generate new unique reagents.", indent=False)
H2("Data and code availability")
for t in ["This paper analyzes existing, publicly available data, accessible at the sources listed in the key resources table. The signal-level filtration-support map generated in this study is provided as Data S1.",
          "All original code has been deposited at Zenodo (https://doi.org/" + CODE_DOI + ") [" + str(CODE_N) + "] and GitHub (" + GITHUB + ") and is publicly available as of the date of publication. The repository contains the analysis scripts unchanged from the versions that produced the reported results, with planned and post hoc analyses distinguished in the script manifest, the analysis parameters, the deviation log, intermediate colocalization results, and Data S1 with its data dictionary. The downstream workflow regenerates the reported results from saved intermediates and was tested on clean copies of the repository; full reconstruction additionally requires the original GWAS, LD and outcome inputs described in the repository and has not been rerun end to end for the released version.",
          "Any additional information required to reanalyze the data reported in this paper is available from the lead contact upon request."]:
    add_rich(d.add_paragraph(style="List Bullet"), t)
H1("Limitations of the study")
for _, el, _ in lim:
    p = P_el(el); p.runs[0].text = "This study has several limitations. " + p.runs[0].text if p.runs else None
H1("Acknowledgments")
P("We thank the investigators of the Kidney Disease Genetic Scorecard, the CKDGen Consortium and FinnGen for making their data publicly available. This research received no specific grant from any funding agency in the public, commercial, or not-for-profit sectors.", indent=False)
H1("Author contributions"); P("H.-L.C.: conceptualization, methodology, software, formal analysis, data curation, visualization, and writing – original draft. H.-H.W.: conceptualization, supervision, project administration, and writing – review & editing. M.-J.J.: supervision and writing – review & editing. Y.-H.T., C.-H.L., and C.-T.T.: validation and writing – review & editing. All authors approved the final manuscript.", indent=False)
H1("Declaration of interests"); P("The authors declare no competing interests.", indent=False)
H1("Declaration of generative AI and AI-assisted technologies in the manuscript preparation process")
P("During the preparation of this work, the authors used Claude (Anthropic) to assist with code development, drafting and language editing, and Perplexity Computer (Perplexity) to assist with manuscript review, reference checking and language revision. The authors reviewed and edited the content as needed and take full responsibility for the publication.", indent=False)
d.add_page_break(); H1("Figure titles and legends")
for kind, el, h3 in S["Figure legends"]:
    if kind == "paragraph" and text_of(el).strip(): P_el(el, indent=False)
d.add_page_break(); H1("Tables"); emit(S["Tables"])
d.add_page_break()
# ---- STAR Methods ----
H1("STAR★Methods")
H2("Key resources table")
KRT = [["REAGENT or RESOURCE", "SOURCE", "IDENTIFIER"],
 ["Deposited data", "", ""],
 ["Kidney Disease Genetic Scorecard and European eGFRcrea GWAS", "Liu et al. [3]", "Figshare 26299093; Table S8 of Liu et al."],
 ["eGFRcys and BUN GWAS summary statistics", "Stanzick et al. [2]", "https://ckdgen.imbi.uni-freiburg.de/datasets"],
 ["FinnGen release 13 summary statistics", "Kurki et al. [20]", "https://www.finngen.fi/en/access_results"],
 ["1000 Genomes Phase 3 European reference", "1000 Genomes Project [9]", "https://ctg.cncr.nl/software/magma (g1000_eur)"],
 ["UK Biobank LD matrices", "Weissbrod et al. [10]", "https://alkesgroup.broadinstitute.org/UKBB_LD/"],
 ["Creatinine-specific and GFR-associated loci", "Titan et al. [5]", "Supplementary Tables S1 and S3 of Titan et al."],
 ["ClinGen gene–disease validity curations", "Strande et al. [15]", "https://search.clinicalgenome.org (accessed 9 Oct 2026)"],
 ["Genomics England PanelApp kidney panels", "Martin et al. [16]", "https://panelapp.genomicsengland.co.uk (accessed 10 Oct 2026)"],
 ["Kidney disease gene list (625 genes)", "Groopman et al. [17]", "Supplementary Appendix of Groopman et al."],
 ["Mouse renal-function phenotypes", "Baldarelli et al. [18]", "https://www.informatics.jax.org (MGI_GenePheno.rpt)"],
 ["Gene annotation (GRCh37)", "Weeks et al. [21]", "https://github.com/FinucaneLab/pops"],
 ["Filtration-support map", "This paper", "Data S1"],
 ["Software and algorithms", "", ""],
 ["R 4.3.3", "R Core Team", "https://www.r-project.org"],
 ["coloc 6.0.3", "Wallace [13]; Giambartolomei et al. [14]", "https://github.com/chr1swallace/coloc"],
 ["susieR 0.12.35", "Wang et al. [11]; Zou et al. [12]", "https://github.com/stephenslab/susieR"],
 ["Python 3.13 (NumPy 2.5.3, pandas 3.0.5, scikit-learn 1.9.1, statsmodels 0.15.0, SciPy 1.18.1, matplotlib 3.11.2)", "Python Software Foundation", "https://www.python.org"],
 ["XGBoost 3.4.1", "Chen and Guestrin [19]", "https://xgboost.ai"],
 ["Analysis code and filtration-support map (version 1.1.0)", "This paper [" + str(CODE_N) + "]", "Zenodo: https://doi.org/" + CODE_DOI + "; GitHub: " + GITHUB]]
SIMPLE_TABLE(KRT)
H2("Experimental model and study participant details")
P("This study used only publicly available summary-level data from previously published genome-wide association studies and the FinnGen release 13 public results; no individual-level data were accessed and no new participants were recruited. Ethics approval and informed consent were obtained by the original studies. All analyzed GWAS summary statistics were of European ancestry, whereas the fine-mapped signal definitions came from the multi-ancestry analysis of Liu et al.; sex- and age-specific analyses were not possible with the summary data used.")
H2("Method details")
emit(S["Methods"], skip_h3=("Software, ethics and data availability",), h3_level=3)
H2("Quantification and statistical analysis")
P("Colocalization posterior probabilities were estimated with coloc-SuSiE, coloc.bf_bf or coloc.abf as described above. Classes were compared with one-sided Wilcoxon rank-sum tests and Hodges\u2013Lehmann shifts with percentile-bootstrap confidence intervals. Gene prioritization was evaluated by pooled held-out average precision, reported as AUPRC, with chromosome-level bootstrap confidence intervals. FinnGen comparisons used one-sided permutation tests (10,000 permutations within propensity-score quintiles), covariate-adjusted logistic regression (with a Firth-penalized sensitivity analysis) and Benjamini\u2013Hochberg correction across the three secondary outcomes. Confidence intervals for mean aligned log odds ratios were the sample mean \u00b1 1.96 standard errors across retained signals, and logistic-regression confidence intervals were Wald intervals on the log-odds scale, exponentiated to odds ratios; these intervals do not incorporate LD among retained variants. Sample sizes (numbers of signals, genes, loci and cases) are given in the Results, tables and figure legends. Analyses used R 4.3.3 and Python 3.13 (key resources table). Numerical differences were calculated before rounding.")

H1("Supplemental information")
P("Document S1. Tables S1 and S2", indent=False)
P("Data S1. Signal-level filtration-support map for 2,051 fine-mapped eGFRcrea signals, with field definitions and recommended use, related to Figure 2", indent=False)
d.add_page_break(); H1("References")
REFS = MASTER_REFS + [CODE_REF]
for n_old in sorted(range(1, len(REFS) + 1), key=lambda n: CITEMAP.get(n, n)):
    p = d.add_paragraph(); p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.first_line_indent = Inches(-0.35)
    p.add_run(f"{CITEMAP.get(n_old, n_old)}. "); add_rich(p, fix(REFS[n_old - 1]))
_seen_before_refs = list(SEEN)
print("references:", len(REFS))
SUPP_REPL = [("Table S1. Deviations from the internally locked analysis plan", "Table S1. Deviations from the internally locked analysis plan, related to STAR Methods"),
             ("Table S2. Covariate balance and propensity-score strata for the FinnGen analysis (dialysis set)", "Table S2. Covariate balance and propensity-score strata for the FinnGen analysis (dialysis set), related to Figure 5")]
if SUBMISSION:
    os.makedirs("manuscript/submission", exist_ok=True)
    d.save("manuscript/submission/Manuscript.docx")
    # Document S1: no title, authors or page numbers (added by the journal)
    new_doc(line_numbers=False, page_numbers=False)
    for sty in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:          # SI guideline: Arial 10 pt
        d.styles[sty].font.name = "Arial"; d.styles[sty].font.size = Pt(10)
        rpr = d.styles[sty].element.get_or_add_rPr(); rf = rpr.get_or_add_rFonts(); rf.set(qn("w:eastAsia"), "Arial")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"): rf.attrib.pop(qn(a), None)
    d.styles["Normal"].paragraph_format.line_spacing = 1.0
    REPL.extend(SUPP_REPL)
    H1("Document S1. Tables S1 and S2"); emit(S["Supplementary tables"])
    d.save("manuscript/submission/Document_S1.docx")
    # Highlights (separate Word file)
    new_doc(line_numbers=False, page_numbers=False)
    H1("Highlights")
    for h in HIGHLIGHTS: d.add_paragraph(h, style="List Bullet")
    d.save("manuscript/submission/Highlights.docx")
    # Key resources table (separate Word file)
    new_doc(line_numbers=False, page_numbers=False)
    H1("Key resources table"); SIMPLE_TABLE(KRT)
    d.save("manuscript/submission/Key_Resources_Table.docx"); print("submission files written")
else:
    d.add_page_break(); H1("Document S1. Supplemental tables"); REPL.extend(SUPP_REPL); emit(S["Supplementary tables"])
    d.add_page_break(); H1("Figures (for review)")
    for f, cap in [("Fig1", "Figure 1"), ("Fig2", "Figure 2"), ("Fig3", "Figure 3"), ("Fig4", "Figure 4"), ("Fig5", "Figure 5"), ("Graphical_abstract", "Graphical abstract")]:
        d.add_paragraph().add_run(cap).bold = True
        src = {"Fig1": "Fig1_study_design", "Fig2": "FigNew_map", "Fig3": "Fig2_fs_titan", "Fig4": "Fig3_incremental_value", "Fig5": "Fig4_finngen"}.get(f)
        d.add_picture(f"figs/{src}.png" if src else "figs/final/Graphical_abstract.png", width=Inches(6.0 if src else 4.0))
    out = "manuscript/iScience_Manuscript_FiltrationSupportMap_draft.docx"; d.save(out); print(out)

if not os.environ.get("CITEMAP"):
    m = {n: i + 1 for i, n in enumerate(_seen_before_refs)}
    missing = [n for n in range(1, len(REFS) + 1) if n not in m]
    for n in missing: m[n] = len(m) + 1
    _json.dump(m, open("/tmp/claude-0/-home-claude/cdd709eb-e6fa-5221-bfa3-4ba0261c0645/scratchpad/citemap.json", "w"))
    print("citation map written; uncited references:", missing)
